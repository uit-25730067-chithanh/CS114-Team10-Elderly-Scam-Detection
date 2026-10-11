#!/usr/bin/env python3
"""
Tự động đóng gói thư mục nộp bài (submit) cho Đồ án CS114 (Nhóm 10)
Chuẩn cấu trúc: CS114_F31_CN2_TTNT - Nhom 10 (Tham chiếu chuẩn CS106 & Đề cương UIT)
"""

import os
import shutil
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

FINAL_DIR = Path(__file__).resolve().parent
DRAFT_DIR = FINAL_DIR / "draft" / "scam-detection"
SUBMIT_ROOT = FINAL_DIR / "submit"
TARGET_DIR = SUBMIT_ROOT / "25730067_25730013_25730061"

EXCLUDE_NAMES = {
    ".git", ".venv", "venv", "__pycache__", ".pytest_cache",
    "node_modules", ".DS_Store", "._*", ".idea", ".vscode", "tmp"
}

def copy_filtered_dir(src: Path, dst: Path, allowed_extensions=None):
    """Sao chép thư mục và loại bỏ các file/thư mục rác."""
    dst.mkdir(parents=True, exist_ok=True)
    for root, dirs, files in os.walk(src):
        # Prune dirs
        dirs[:] = [d for d in dirs if d not in EXCLUDE_NAMES and not d.startswith(".")]
        rel_path = Path(root).relative_to(src)
        dest_dir = dst / rel_path
        dest_dir.mkdir(parents=True, exist_ok=True)

        for f in files:
            if f in EXCLUDE_NAMES or f.startswith(".") or f.startswith("~$"):
                continue
            if allowed_extensions and Path(f).suffix.lower() not in allowed_extensions:
                continue
            src_file = Path(root) / f
            dest_file = dest_dir / f
            shutil.copy2(src_file, dest_file)

def create_danh_sach_nhom_excel(output_path: Path):
    """Tạo bảng Excel danh sách nhóm và phân công nhiệm vụ chuẩn UIT."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh sách nhóm"

    # Định dạng
    font_title = Font(name="Times New Roman", size=14, bold=True, color="1F497D")
    font_subtitle = Font(name="Times New Roman", size=11, italic=True, color="595959")
    font_header = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
    font_content = Font(name="Times New Roman", size=11)
    font_content_bold = Font(name="Times New Roman", size=11, bold=True)

    fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    fill_even = PatternFill(start_color="F2F5F8", end_color="F2F5F8", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    # Tiêu đề
    ws.merge_cells("A1:H1")
    ws["A1"] = "BẢNG DANH SÁCH NHÓM VÀ PHÂN CÔNG NHIỆM VỤ ĐỒ ÁN CUỐI KỲ"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:H2")
    ws["A2"] = "Môn học: Máy học (CS114.F31.CN2.TTNT) — Giảng viên hướng dẫn: ThS. Cáp Phạm Đình Thăng"
    ws["A2"].font = font_subtitle
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A3:H3")
    ws["A3"] = "Đề tài: Phát hiện Tin nhắn & Kịch bản Thao túng Tâm lý Lừa đảo nhắm vào Người cao tuổi tại Việt Nam — Nhóm 10"
    ws["A3"].font = font_subtitle
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")

    # Header bảng
    headers = [
        "STT", "Họ và Tên", "MSSV", "Lớp", "Vai trò",
        "Nhiệm vụ chính phụ trách", "Tỷ lệ hoàn thành", "Đóng góp"
    ]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Dữ liệu 3 thành viên
    members = [
        (
            1, "Đặng Chí Thanh", "25730067", "CS114.F31.CN2.TTNT", "Trưởng nhóm",
            "Quản lý tiến độ đồ án; Thiết kế kiến trúc tổng thể; Tiền xử lý NLP tiếng Việt & Trích xuất đặc trưng; Đánh giá mô hình & Giải thích SHAP; Soạn thảo Báo cáo học thuật (Word/PDF).",
            "100%", "33.3%"
        ),
        (
            2, "Bùi Thị Mỷ Cẩm", "25730013", "CS114.F31.CN2.TTNT", "Thành viên",
            "Thu thập & kiểm định dữ liệu SMS lừa đảo; Xây dựng, huấn luyện và tối ưu hóa các mô hình Naive Bayes, XGBoost, Linear SVM; Thiết kế Slide báo cáo thuyết trình (PPTX/PDF).",
            "100%", "33.3%"
        ),
        (
            3, "Hoàng Cao Sơn", "25730061", "CS114.F31.CN2.TTNT", "Thành viên",
            "Thu thập dữ liệu phóng sự báo chí & truyền hình; Huấn luyện mô hình Random Forest; Phát triển giao diện Web Demo Streamlit tương tác; Xây dựng kịch bản test demo.",
            "100%", "33.3%"
        ),
    ]

    for row_idx, member in enumerate(members, start=6):
        is_even = (row_idx % 2 == 0)
        for col_idx, val in enumerate(member, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_content_bold if col_idx in (2, 5) else font_content
            if is_even:
                cell.fill = fill_even
            cell.border = thin_border
            align_h = "center" if col_idx in (1, 3, 4, 5, 7, 8) else "left"
            cell.alignment = Alignment(horizontal=align_h, vertical="center", wrap_text=True)

    # Chỉnh độ rộng cột
    col_widths = {1: 8, 2: 24, 3: 14, 4: 24, 5: 16, 6: 55, 7: 18, 8: 14}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 20
    ws.row_dimensions[5].height = 26
    for r in range(6, 9):
        ws.row_dimensions[r].height = 50

    wb.save(str(output_path))
    print(f"[✓] Đã tạo file Danh_sach_nhom.xlsx: {output_path}")

def create_chay_demo_bat(output_path: Path):
    """Tạo file Chay_Demo.bat tiện lợi cho giảng viên."""
    content = """@echo off
chcp 65001 >nul
set PYTHONIOENCODING=utf-8
title He Thong Phat Hien Tin Nhan Lua Dao - CS114 Nhom 10

echo ======================================================================
echo    HỆ THỐNG PHÁT HIỆN TIN NHẮN LỪA ĐẢO CHO NGƯỜI CAO TUỔI (CS114)
echo    Nhóm 10 - Lớp CS114.F31.CN2.TTNT
echo ======================================================================
echo.
echo [*] Đang khởi động ứng dụng Web Demo tương tác (Streamlit)...
echo [*] Trình duyệt web sẽ tự động mở trang ứng dụng ngay bây giờ...
echo.

cd /d "%~dp0demo"
python -m streamlit run app.py

pause
"""
    output_path.write_text(content, encoding="utf-8")
    print(f"[✓] Đã tạo file Chay_Demo.bat: {output_path}")

def create_huong_dan_su_dung_md(output_path: Path):
    """Tạo tài liệu Hướng dẫn sử dụng Markdown chi tiết."""
    content = """# HƯỚNG DẪN CÀI ĐẶT VÀ SỬ DỤNG CHƯƠNG TRÌNH
## Đồ án môn học: Máy học (CS114.F31.CN2.TTNT) — Nhóm 10
**Đề tài:** Nghiên cứu và Xây dựng Hệ thống Phát hiện Tin nhắn & Kịch bản Thao túng Tâm lý Lừa đảo nhắm vào Người cao tuổi tại Việt Nam  
**Giảng viên hướng dẫn:** ThS. Cáp Phạm Đình Thăng  
**Sinh viên thực hiện:**
1. Đặng Chí Thanh (MSSV: 25730067) — Trưởng nhóm
2. Bùi Thị Mỷ Cẩm (MSSV: 25730013)
3. Hoàng Cao Sơn (MSSV: 25730061)

---

## 1. YÊU CẦU MÔI TRƯỜNG
- **Hệ điều hành:** Windows 10/11, macOS, hoặc Linux (Ubuntu 20.04+).
- **Python:** Phiên bản Python 3.9 đến 3.11 (khuyến nghị Python 3.10 hoặc 3.11).
- **Trình quản lý gói:** `pip` mới nhất.

---

## 2. HƯỚNG DẪN CÀI ĐẶT

### Bước 1: Mở Terminal / Command Prompt tại thư mục `Chuong_trinh/`
```bash
cd Chuong_trinh
```

### Bước 2: Tạo và kích hoạt môi trường ảo (Khuyến nghị)
- **Trên Windows:**
```cmd
python -m venv .venv
.venv\\Scripts\\activate
```

- **Trên macOS / Linux:**
```bash
python3 -m venv .venv
source .venv/bin/activate
```

### Bước 3: Cài đặt các thư viện phụ thuộc
```bash
pip install --upgrade pip
pip install -r requirements.txt
```

---

## 3. HƯỚNG DẪN CHẠY ỨNG DỤNG WEB DEMO

### Cách 1: Chạy nhanh trên Windows bằng file batch
- Nhấp đúp chuột vào file `Chay_Demo.bat` trong thư mục `Chuong_trinh/`.
- Ứng dụng sẽ tự động kích hoạt và mở giao diện trình duyệt web tại địa chỉ `http://localhost:8501`.

### Cách 2: Chạy thủ công qua dòng lệnh (Windows / macOS / Linux)
```bash
cd demo
python -m streamlit run app.py
```

### Các tính năng chính trên Web Demo:
1. **Phân tích tin nhắn theo thời gian thực:** Nhập nội dung tin nhắn tiếng Việt bất kỳ, hệ thống sẽ tiền xử lý, trích xuất đặc trưng và đưa ra cảnh báo An toàn / Nguy cơ Lừa đảo.
2. **Đối sánh đa mô hình:** Lựa chọn và chuyển đổi linh hoạt giữa 4 mô hình: XGBoost (Mặc định), Linear SVM, Random Forest, Naive Bayes.
3. **Giải thích trực quan (XAI):** Hiển thị thang đo rủi ro (Risk Meter), từ khóa thao túng tâm lý được tô màu, và lý do dự đoán theo mô hình SHAP.
4. **Kho mẫu thử nghiệm thực tế:** Nạp nhanh các kịch bản lừa đảo phổ biến (mạo danh bệnh viện tặng yến sào, mạo danh công an phạt nguội, việc làm xem video, thần y chữa khớp...).

---

## 4. HƯỚNG DẪN CHẠY CÁC JUPYTER NOTEBOOKS

Toàn bộ quá trình nghiên cứu, khám phá dữ liệu, huấn luyện mô hình và đánh giá giải thích được tổ chức theo thứ tự logic trong thư mục `code/notebooks/`:

```
code/notebooks/
├── 01_preprocessing.ipynb          # Tiền xử lý văn bản tiếng Việt & tách từ (PyVi)
├── 02_eda.ipynb                    # Khám phá dữ liệu (EDA) & phân tích phân phối
├── 03_model_naive_bayes.ipynb      # Huấn luyện mô hình cơ sở Multinomial Naive Bayes
├── 04_model_random_forest.ipynb    # Huấn luyện mô hình Random Forest
├── 04_model_xgboost.ipynb          # Huấn luyện & tối ưu mô hình XGBoost (scale_pos_weight)
├── 05_model_svm.ipynb              # Huấn luyện mô hình Linear SVM & phân tích siêu phẳng
└── 06_evaluation_xai.ipynb         # Đánh giá đối sánh chéo đa mô hình & giải thích SHAP
```

Để mở và thực thi các notebook:
```bash
cd code
jupyter notebook notebooks/
```

---

## 5. CẤU TRÚC MÃ NGUỒN (`code/src/`)
- `src/preprocessing/`: Các module chuẩn hóa Unicode, khử PII, tách từ tiếng Việt.
- `src/models/`: Cấu hình siêu tham số và huấn luyện 4 mô hình.
- `src/evaluation/`: Các hàm đo đạc chỉ số (Accuracy, Precision, Recall, F1, ROC-AUC) và tính ma trận nhầm lẫn.
- `src/utils/`: Thiết lập seed ngẫu nhiên (`RANDOM_STATE=42`) đảm bảo tính tái lập 100%.
"""
    output_path.write_text(content, encoding="utf-8")
    print(f"[✓] Đã tạo file Huong_dan_su_dung.md: {output_path}")

def main():
    print(f"[*] Bắt đầu chuẩn bị đóng gói folder submit...")
    
    # 1. Khởi tạo thư mục đích
    if TARGET_DIR.exists():
        shutil.rmtree(TARGET_DIR)
    
    bao_cao_dir = TARGET_DIR / "Bao_cao"
    chuong_trinh_dir = TARGET_DIR / "Chuong_trinh"
    code_dir = chuong_trinh_dir / "code"
    demo_dir = chuong_trinh_dir / "demo"

    bao_cao_dir.mkdir(parents=True, exist_ok=True)
    code_dir.mkdir(parents=True, exist_ok=True)
    demo_dir.mkdir(parents=True, exist_ok=True)

    # 2. Đóng gói Bao_cao/
    print("[*] Đang sao chép Báo cáo & Slide vào Bao_cao/...")
    src_docx = DRAFT_DIR / "reports" / "report" / "report.docx"
    src_pdf = DRAFT_DIR / "reports" / "report" / "report.pdf"
    src_pptx = DRAFT_DIR / "reports" / "slides" / "phase09_presentation.pptx"
    src_slide_pdf = DRAFT_DIR / "reports" / "slides" / "phase09_presentation.pdf"

    if src_docx.exists():
        shutil.copy2(src_docx, bao_cao_dir / "[Nhom10]_Report_ScamDetection.docx")
    if src_pdf.exists():
        shutil.copy2(src_pdf, bao_cao_dir / "[Nhom10]_Report_ScamDetection.pdf")
    if src_pptx.exists():
        shutil.copy2(src_pptx, bao_cao_dir / "[Nhom10]_Slide_ScamDetection.pptx")
    if src_slide_pdf.exists():
        shutil.copy2(src_slide_pdf, bao_cao_dir / "[Nhom10]_Slide_ScamDetection.pdf")

    # 3. Đóng gói Chuong_trinh/code/
    print("[*] Đang sao chép mã nguồn kỹ thuật vào Chuong_trinh/code/...")
    # src
    copy_filtered_dir(DRAFT_DIR / "src", code_dir / "src")
    
    # notebooks (chỉ lấy ipynb và README)
    copy_filtered_dir(DRAFT_DIR / "notebooks", code_dir / "notebooks", allowed_extensions={".ipynb", ".md"})
    
    # data
    data_raw_dst = code_dir / "data" / "raw"
    data_proc_dst = code_dir / "data" / "processed"
    data_raw_dst.mkdir(parents=True, exist_ok=True)
    data_proc_dst.mkdir(parents=True, exist_ok=True)
    
    # raw: copy dataset chính và template
    for f in ["dataset.csv", "dataset_template.csv", "README.md"]:
        p = DRAFT_DIR / "data" / "raw" / f
        if p.exists():
            shutil.copy2(p, data_raw_dst / f)
            
    # processed: copy features, scalers, vectorizers
    copy_filtered_dir(DRAFT_DIR / "data" / "processed", data_proc_dst)

    # models: copy 4 mô hình + best bundle + metrics json
    models_dst = code_dir / "models"
    models_dst.mkdir(parents=True, exist_ok=True)
    copy_filtered_dir(DRAFT_DIR / "models", models_dst)

    # reports: figures và model_comparison.csv
    reports_dst = code_dir / "reports"
    reports_dst.mkdir(parents=True, exist_ok=True)
    if (DRAFT_DIR / "reports" / "figures").exists():
        copy_filtered_dir(DRAFT_DIR / "reports" / "figures", reports_dst / "figures")
    if (DRAFT_DIR / "reports" / "model_comparison.csv").exists():
        shutil.copy2(DRAFT_DIR / "reports" / "model_comparison.csv", reports_dst / "model_comparison.csv")

    # 4. Đóng gói Chuong_trinh/demo/
    print("[*] Đang sao chép giao diện Web Demo vào Chuong_trinh/demo/...")
    copy_filtered_dir(DRAFT_DIR / "demo", demo_dir)

    # 5. Các tệp điều hành trong Chuong_trinh/
    create_chay_demo_bat(chuong_trinh_dir / "Chay_Demo.bat")
    create_huong_dan_su_dung_md(chuong_trinh_dir / "Huong_dan_su_dung.md")
    
    # requirements.txt
    req_src = DRAFT_DIR / "requirements.txt"
    if req_src.exists():
        shutil.copy2(req_src, chuong_trinh_dir / "requirements.txt")
        shutil.copy2(req_src, code_dir / "requirements.txt")

    # 6. Các tệp cấp gốc submit/
    print("[*] Đang khởi tạo Danh_sach_nhom.xlsx và Kịch bản test demo...")
    create_danh_sach_nhom_excel(TARGET_DIR / "Danh_sach_nhom.xlsx")

    test_script_src = FINAL_DIR / "kich_ban_test_demo_nhom10.xlsx"
    if test_script_src.exists():
        shutil.copy2(test_script_src, TARGET_DIR / "kich_ban_test_demo_nhom10.xlsx")

    print(f"\n[✓] HOÀN TẤT ĐÓNG GÓI THƯ MỤC NỘP BÀI:")
    print(f"    Target: {TARGET_DIR}")

if __name__ == "__main__":
    main()
