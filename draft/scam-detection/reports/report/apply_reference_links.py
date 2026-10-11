import sys
import shutil
import re
import os
import docx
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def main():
    report_path = sys.argv[1] if len(sys.argv) > 1 else "draft/scam-detection/reports/report/report.docx"
    backup_path = report_path.replace(".docx", "_backup_before_links.docx")

    if not os.path.exists(report_path):
        print(f"Error: {report_path} not found!")
        return

    # Backup first
    shutil.copy2(report_path, backup_path)
    print(f"Backed up to {backup_path}")

    doc = docx.Document(report_path)
    print(f"Total paragraphs in document: {len(doc.paragraphs)}")

    url_regex = re.compile(r'(https?://[^\s]+)')
    doi_regex = re.compile(r'doi:\s*([^\s]+)')

    modified_count = 0

    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if not re.match(r'^\[\d+\]', text):
            continue

        url_m = url_regex.search(text)
        doi_m = doi_regex.search(text)

        if url_m or doi_m:
            if url_m:
                raw_url = url_m.group(1).rstrip('.')
                target = raw_url
                start, end = url_m.span(1)
                if url_m.group(1).endswith('.'):
                    end -= 1
            else:
                raw_doi = doi_m.group(1).rstrip('.')
                target = f"https://doi.org/{raw_doi}"
                start = doi_m.start(1)
                end = start + len(raw_doi)

            before = text[:start]
            link_str = text[start:end]
            after = text[end:]

            assert before + link_str + after == text, f"Mismatch in P#{i}"

            # Create external relationship for hyperlink
            r_id = p.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

            # Preserve pPr (hanging indent, line spacing, etc.)
            pPr = p._p.pPr

            # Remove existing run children (preserve pPr and bookmarks)
            for child in list(p._p):
                if child != pPr and not child.tag.endswith(('bookmarkStart', 'bookmarkEnd')):
                    p._p.remove(child)

            # 1. Before run
            if before:
                before_escaped = before.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                r_before = parse_xml(
                    f'<w:r {nsdecls("w")}>'
                    f'<w:rPr>'
                    f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>'
                    f'<w:sz w:val="22"/>'
                    f'<w:color w:val="000000"/>'
                    f'</w:rPr>'
                    f'<w:t xml:space="preserve">{before_escaped}</w:t>'
                    f'</w:r>'
                )
                p._p.append(r_before)

            # 2. Hyperlink element
            link_escaped = link_str.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            h_xml = (
                f'<w:hyperlink {nsdecls("w", "r")} r:id="{r_id}" w:history="1">'
                f'<w:r>'
                f'<w:rPr>'
                f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>'
                f'<w:sz w:val="22"/>'
                f'<w:color w:val="0563C1"/>'
                f'<w:u w:val="single"/>'
                f'</w:rPr>'
                f'<w:t xml:space="preserve">{link_escaped}</w:t>'
                f'</w:r>'
                f'</w:hyperlink>'
            )
            h_elem = parse_xml(h_xml)
            p._p.append(h_elem)

            # 3. After run
            if after:
                after_escaped = after.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                r_after = parse_xml(
                    f'<w:r {nsdecls("w")}>'
                    f'<w:rPr>'
                    f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>'
                    f'<w:sz w:val="22"/>'
                    f'<w:color w:val="000000"/>'
                    f'</w:rPr>'
                    f'<w:t xml:space="preserve">{after_escaped}</w:t>'
                    f'</w:r>'
                )
                p._p.append(r_after)

            modified_count += 1
            print(f"Applied link [{modified_count}]: P#{i} -> {target}")

    doc.save(report_path)
    print(f"Saved {report_path} successfully! Total links applied: {modified_count}")

    # Validate
    doc_check = docx.Document(report_path)
    print(f"Validation: loaded document with {len(doc_check.paragraphs)} paragraphs.")
    ref_checks = [p.text for p in doc_check.paragraphs if re.match(r'^\[\d+\]', p.text.strip())]
    print(f"Found {len(ref_checks)} reference entries verified.")

if __name__ == "__main__":
    main()
