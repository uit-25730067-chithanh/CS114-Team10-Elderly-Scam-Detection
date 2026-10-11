import sys
import os
import re
import docx
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def apply_citation_links(docx_path):
    print(f"[*] Đang tải tệp tài liệu: {docx_path}")
    if not os.path.exists(docx_path):
        print(f"[!] Không tìm thấy tệp: {docx_path}")
        return False

    doc = docx.Document(docx_path)
    
    # ══════════════════════════════════════════════════════════════════════════
    # BƯỚC 1: ĐẶT 20 BOOKMARKS CHO DANH MỤC TÀI LIỆU THAM KHẢO
    # ══════════════════════════════════════════════════════════════════════════
    ref_paragraphs = {}
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        m = re.match(r'^\[(\d+)\]', text)
        if m:
            num = int(m.group(1))
            if 1 <= num <= 20:
                ref_paragraphs[num] = (i, p)

    print(f"[*] Tìm thấy {len(ref_paragraphs)}/20 mục tài liệu tham khảo ở cuối bài.")
    
    bookmarks_added = 0
    for num in sorted(ref_paragraphs.keys()):
        p_idx, p = ref_paragraphs[num]
        bm_name = f"_Ref_Item_{num:02d}"
        bm_id = 500 + num
        
        # Kiểm tra xem bookmark đã tồn tại chưa
        existing = p._p.xpath(f'.//w:bookmarkStart[@w:name="{bm_name}"]')
        if not existing:
            bm_start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{bm_id}" w:name="{bm_name}"/>')
            bm_end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{bm_id}"/>')
            
            # Chèn vào ngay sau pPr (nếu có)
            pPr = p._p.pPr
            if pPr is not None:
                pPr.addnext(bm_end)
                pPr.addnext(bm_start)
            else:
                p._p.insert(0, bm_end)
                p._p.insert(0, bm_start)
            bookmarks_added += 1
            print(f"  [✓] Đã tạo Bookmark '{bm_name}' tại Ref #[{num}] (P#{p_idx})")

    print(f"[*] Tổng số Bookmark tài liệu tham khảo đã tạo: {bookmarks_added}")

    # ══════════════════════════════════════════════════════════════════════════
    # BƯỚC 2: QUÉT VÀ THAY THẾ IN-TEXT CITATIONS THÀNH HYPERLINKS NỘI BỘ
    # ══════════════════════════════════════════════════════════════════════════
    cite_regex = re.compile(r'\[(\d+)\]')
    citation_links_count = 0

    def process_paragraph_runs(p, is_table_cell=False):
        nonlocal citation_links_count
        p_text = p.text.strip()
        # Bỏ qua nếu là đoạn tài liệu tham khảo chính
        if re.match(r'^\[\d+\]', p_text):
            return

        # Lấy danh sách các run con trực tiếp của p._p (không nằm trong hyperlink sẵn có)
        child_elements = list(p._p)
        for child in child_elements:
            if child.tag.endswith('}r'):
                # Lấy text của run
                t_elements = child.findall(qn('w:t'))
                if not t_elements:
                    continue
                r_text = "".join([t.text for t in t_elements if t.text is not None])
                if not r_text:
                    continue

                matches = list(cite_regex.finditer(r_text))
                valid_matches = [m for m in matches if 1 <= int(m.group(1)) <= 20]
                if not valid_matches:
                    continue

                # Có citation cần bọc hyperlink: tách run này ra
                rPr = child.find(qn('w:rPr'))
                # Lấy font name và size để giữ nguyên
                font_name = "Times New Roman"
                font_size = "26" # 13pt mặc định (26 half-points)
                if is_table_cell:
                    font_size = "17" # 8.5pt mặc định trong bảng
                
                if rPr is not None:
                    sz_elem = rPr.find(qn('w:sz'))
                    if sz_elem is not None and sz_elem.get(qn('w:val')):
                        font_size = sz_elem.get(qn('w:val'))
                    rf_elem = rPr.find(qn('w:rFonts'))
                    if rf_elem is not None and rf_elem.get(qn('w:ascii')):
                        font_name = rf_elem.get(qn('w:ascii'))

                # Tạo danh sách các phần tử thay thế
                new_elements = []
                last_pos = 0

                for m in valid_matches:
                    start_pos, end_pos = m.span()
                    num = int(m.group(1))
                    bm_name = f"_Ref_Item_{num:02d}"

                    # 1. Đoạn text trước match
                    if start_pos > last_pos:
                        before_text = r_text[last_pos:start_pos]
                        before_esc = before_text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                        r_before = parse_xml(
                            f'<w:r {nsdecls("w")}>'
                            f'<w:rPr>'
                            f'<w:rFonts w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>'
                            f'<w:sz w:val="{font_size}"/>'
                            f'</w:rPr>'
                            f'<w:t xml:space="preserve">{before_esc}</w:t>'
                            f'</w:r>'
                        )
                        new_elements.append(r_before)

                    # 2. Hyperlink bọc citation "[num]"
                    cite_text = f"[{num}]"
                    cite_hl = parse_xml(
                        f'<w:hyperlink {nsdecls("w")} w:anchor="{bm_name}" w:history="1">'
                        f'<w:r>'
                        f'<w:rPr>'
                        f'<w:rFonts w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>'
                        f'<w:sz w:val="{font_size}"/>'
                        f'<w:color w:val="0563C1"/>'
                        f'<w:u w:val="none"/>'
                        f'</w:rPr>'
                        f'<w:t xml:space="preserve">{cite_text}</w:t>'
                        f'</w:r>'
                        f'</w:hyperlink>'
                    )
                    new_elements.append(cite_hl)
                    citation_links_count += 1
                    last_pos = end_pos

                # 3. Đoạn text còn lại sau match cuối cùng
                if last_pos < len(r_text):
                    after_text = r_text[last_pos:]
                    after_esc = after_text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                    r_after = parse_xml(
                        f'<w:r {nsdecls("w")}>'
                        f'<w:rPr>'
                        f'<w:rFonts w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>'
                        f'<w:sz w:val="{font_size}"/>'
                        f'</w:rPr>'
                        f'<w:t xml:space="preserve">{after_esc}</w:t>'
                        f'</w:r>'
                    )
                    new_elements.append(r_after)

                # Thay thế run cũ bằng danh sách new_elements
                idx = p._p.index(child)
                p._p.remove(child)
                for offset, elem in enumerate(new_elements):
                    p._p.insert(idx + offset, elem)

    # Quét toàn bộ paragraphs thông thường
    for p in doc.paragraphs:
        process_paragraph_runs(p, is_table_cell=False)

    # Quét toàn bộ paragraphs trong tables (trừ Table 1 là Mục lục)
    for t_idx, table in enumerate(doc.tables):
        if t_idx == 1: # Bỏ qua Mục lục
            continue
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    process_paragraph_runs(p, is_table_cell=True)

    print(f"[*] Đã liên kết thành công {citation_links_count} in-text citations trỏ về danh mục Tài liệu tham khảo!")
    
    doc.save(docx_path)
    print(f"[✓] Đã lưu thành công tệp: {docx_path}")
    return True

if __name__ == "__main__":
    report_file = sys.argv[1] if len(sys.argv) > 1 else "draft/scam-detection/reports/report/report.docx"
    apply_citation_links(report_file)
