import sys
import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def apply_interactive_toc(docx_path):
    print(f"[*] Đang tải tệp: {docx_path}")
    doc = docx.Document(docx_path)
    
    if len(doc.tables) < 2:
        print("[!] Không tìm thấy bảng Mục lục (Table 1) trong tài liệu!")
        return False
        
    toc_tbl = doc.tables[1]
    toc_rows = toc_tbl.rows[1:] # Bỏ hàng tiêu đề CẤU TRÚC BÁO CÁO ĐỒ ÁN
    print(f"[*] Tìm thấy {len(toc_rows)} mục trong bảng Mục lục.")
    
    bm_counter = 100
    success_count = 0
    
    for idx, r in enumerate(toc_rows):
        title = r.cells[0].text.strip()
        t_clean = title.replace('    ', '').strip()
        page_str = r.cells[1].text.strip()
        
        # 1. Tìm paragraph đích
        target_p = None
        target_p_idx = None
        
        # Thử tìm chính xác trước
        for p_idx, p in enumerate(doc.paragraphs):
            p_text = p.text.strip()
            if not p_text:
                continue
            if t_clean == p_text:
                target_p = p
                target_p_idx = p_idx
                break
                
        # Nếu chưa thấy, tìm theo chuỗi con (với độ dài >= 8 ký tự để tránh nhầm)
        if target_p is None and len(t_clean) >= 8:
            for p_idx, p in enumerate(doc.paragraphs):
                p_text = p.text.strip()
                if not p_text:
                    continue
                if t_clean in p_text or p_text in t_clean:
                    target_p = p
                    target_p_idx = p_idx
                    break
                    
        # Xử lý ngoại lệ đặc thù
        if target_p is None:
            if '4.2.2' in t_clean and 'ROC' in t_clean:
                for p_idx, p in enumerate(doc.paragraphs):
                    if '4.2.2' in p.text and 'ROC' in p.text:
                        target_p = p
                        target_p_idx = p_idx
                        break
            elif 'TÀI LIỆU THAM KHẢO' in t_clean:
                for p_idx, p in enumerate(doc.paragraphs):
                    if 'TÀI LIỆU THAM KHẢO' in p.text:
                        target_p = p
                        target_p_idx = p_idx
                        break
                        
        if target_p is None:
            print(f"[!] Không tìm thấy đoạn văn tương ứng cho mục #{idx}: '{title}'")
            continue
            
        # 2. Tạo Bookmark duy nhất tại target_p
        bm_counter += 1
        bm_id = str(bm_counter)
        bm_name = f"_Toc_Sec_{idx+1:02d}"
        
        bm_start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{bm_id}" w:name="{bm_name}"/>')
        bm_end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{bm_id}"/>')
        
        # Chèn bm_start ngay sau pPr (nếu có) hoặc đầu paragraph
        if target_p._p.pPr is not None:
            target_p._p.insert(target_p._p.index(target_p._p.pPr) + 1, bm_start)
        else:
            target_p._p.insert(0, bm_start)
        target_p._p.append(bm_end)
        
        # 3. Biến text ô Cột 0 (Tên mục) thành Hyperlink trỏ về Bookmark
        cell0_p = r.cells[0].paragraphs[0]
        runs0 = [c for c in cell0_p._p if c.tag.endswith('}r')]
        if runs0:
            hl0 = parse_xml(f'<w:hyperlink {nsdecls("w")} w:anchor="{bm_name}"/>')
            for run_elem in runs0:
                cell0_p._p.remove(run_elem)
                hl0.append(run_elem)
            cell0_p._p.append(hl0)
            
        # 4. Biến text ô Cột 1 (Số trang) thành Hyperlink trỏ về Bookmark
        cell1_p = r.cells[1].paragraphs[0]
        runs1 = [c for c in cell1_p._p if c.tag.endswith('}r')]
        if runs1:
            hl1 = parse_xml(f'<w:hyperlink {nsdecls("w")} w:anchor="{bm_name}"/>')
            for run_elem in runs1:
                cell1_p._p.remove(run_elem)
                hl1.append(run_elem)
            cell1_p._p.append(hl1)
            
        success_count += 1
        print(f"  [✓] Đã liên kết mục #{idx+1:02d}: '{t_clean[:32]}...' -> Bookmark '{bm_name}' (P#{target_p_idx})")
        
    print(f"\n[*] Đã liên kết thành công {success_count}/{len(toc_rows)} mục trong Mục lục!")
    doc.save(docx_path)
    print(f"[✓] Đã lưu thành công vào tệp: {docx_path}")
    return True

if __name__ == "__main__":
    report_file = sys.argv[1] if len(sys.argv) > 1 else "draft/scam-detection/reports/report/report.docx"
    apply_interactive_toc(report_file)
